import argparse
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_BREAK, WD_LINE_SPACING, WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "resume" / "梁磊-Android端侧AI-AIOS架构师.md"

FONT_CN = "微软雅黑"
FONT_EN = "Aptos"
ACCENT = RGBColor(31, 78, 121)
TEXT = RGBColor(35, 35, 35)
MUTED = RGBColor(85, 85, 85)


def set_cell_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "1F4E79")
    borders.append(bottom)
    p_pr.append(borders)


def apply_run_font(run, size=9.5, bold=False, color=TEXT):
    run.font.name = FONT_EN
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_inline(paragraph, text, size=9.5, color=TEXT):
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        content = part[2:-2] if bold else part
        run = paragraph.add_run(content)
        apply_run_font(run, size=size, bold=bold, color=color)


def paragraph_format(paragraph, before=0, after=2, line=1.05):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
    fmt.line_spacing = line


def add_title(document, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    paragraph_format(paragraph, after=3)
    run = paragraph.add_run(text)
    apply_run_font(run, size=20, bold=True, color=ACCENT)


def add_heading(document, text, level):
    paragraph = document.add_paragraph()
    if level == 2:
        paragraph_format(paragraph, before=7, after=3)
        run = paragraph.add_run(text)
        apply_run_font(run, size=12, bold=True, color=ACCENT)
        set_cell_border(paragraph)
    else:
        paragraph_format(paragraph, before=5, after=1)
        run = paragraph.add_run(text)
        apply_run_font(run, size=10.5, bold=True, color=TEXT)
    return paragraph


def add_body(document, text):
    paragraph = document.add_paragraph()
    paragraph_format(paragraph, after=2)
    add_inline(paragraph, text)


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph_format(paragraph, after=1.5)
    paragraph.paragraph_format.left_indent = Cm(0.45)
    paragraph.paragraph_format.first_line_indent = Cm(-0.2)
    add_inline(paragraph, text)


def build_document(source, output):
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.55)
    section.right_margin = Cm(1.55)

    normal = document.styles["Normal"]
    normal.font.name = FONT_EN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    normal.font.size = Pt(9.5)

    bullet_style = document.styles["List Bullet"]
    bullet_style.font.name = FONT_EN
    bullet_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    bullet_style.font.size = Pt(9.5)

    lines = source.read_text(encoding="utf-8").splitlines()
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("# "):
            add_title(document, stripped[2:])
        elif stripped.startswith("## "):
            add_heading(document, stripped[3:], 2)
        elif stripped.startswith("### "):
            add_heading(document, stripped[4:], 3)
        elif stripped.startswith("- "):
            add_bullet(document, stripped[2:])
        else:
            paragraph = document.add_paragraph()
            if stripped.startswith(("**Android 端侧", "**车载 Android")):
                paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                paragraph_format(paragraph, after=2)
                add_inline(paragraph, stripped, size=11, color=ACCENT)
            elif stripped.startswith("北京｜"):
                paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                paragraph_format(paragraph, after=1)
                add_inline(paragraph, stripped, size=9, color=MUTED)
            elif stripped.startswith("电话："):
                paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                paragraph_format(paragraph, after=4)
                add_inline(paragraph, stripped, size=9, color=MUTED)
            else:
                paragraph_format(paragraph, after=2)
                add_inline(paragraph, stripped)

    core_properties = document.core_properties
    core_properties.title = source.stem
    core_properties.subject = "求职简历"
    core_properties.author = "梁磊"

    document.save(output)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generate a formatted resume DOCX from Markdown.")
    parser.add_argument("source", nargs="?", type=Path, default=DEFAULT_SOURCE)
    args = parser.parse_args()
    source = args.source.resolve()
    output = source.with_suffix(".docx")
    build_document(source, output)
